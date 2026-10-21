"""考题一：Excel 报表生成脚本（分模块实现）。

模块划分：
    M1  读取 + 清洗      -> read_raw / clean_data
    M2  统计聚合          -> build_salesman_summary / build_region_summary / build_monthly_summary
    M3  Excel 输出        -> write_report
    M4  异常与 CLI        -> generate_report / main
"""
from __future__ import annotations

import argparse
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pandas as pd

# ---- 需求澄清结论（见 需求分析.md 第七节）-------------------------------
# 关键字段 = 全部 6 个字段，任一缺失即删除该行
REQUIRED_COLUMNS = ["日期", "销售员", "产品", "数量", "单价", "地区"]
# 需要做“去空白 + 空字符串视为缺失”的文本列
TEXT_COLUMNS = ["销售员", "产品", "地区"]
AMOUNT_COLUMN = "销售额"
AMOUNT_DECIMALS = 2


class ReportError(Exception):
    """报表生成过程中的可预期错误（文件缺失、字段缺失、数据为空等）。"""


class MissingColumnsError(ReportError):
    """原始表缺少必需字段列。"""


class NoValidDataError(ReportError):
    """清洗后没有任何有效数据。"""


# ======================= M1：读取 + 清洗 =================================

def read_raw(path: str | Path) -> pd.DataFrame:
    """以只读方式读取原始 Excel，并校验必需字段列。"""
    path = Path(path)
    if not path.exists():
        raise ReportError(f"原始文件不存在：{path}")
    if path.suffix.lower() != ".xlsx":
        raise ReportError(f"原始文件必须是 .xlsx 格式：{path.name}")
    try:
        df = pd.read_excel(path, engine="openpyxl")
    except Exception as exc:  # noqa: BLE001 - 统一转成可读错误
        raise ReportError(f"读取 Excel 失败：{path.name}（{exc}）") from exc

    if df.empty:
        raise ReportError(f"原始表为空（仅有表头，无数据行）：{path.name}")

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise MissingColumnsError(
            f"原始表缺少必需字段列：{'、'.join(missing)}；"
            f"实际列为：{'、'.join(map(str, df.columns))}"
        )
    return df


def clean_data(df: pd.DataFrame) -> pd.DataFrame:
    """按拍板规则清洗：删关键字段缺失行 -> 剔除数量<=0 -> 按全字段去重。

    不修改入参 DataFrame（保证原始数据只读）。
    """
    out = df.copy()

    # 文本列规范化：去首尾空白，空字符串视为缺失
    for col in TEXT_COLUMNS:
        if col in out.columns:
            out[col] = out[col].astype("string").str.strip()
            out[col] = out[col].mask(out[col] == "", pd.NA)

    # 数值/日期列转换：无法转换的值 -> 缺失（后续按关键字段缺失删除）
    out["日期"] = pd.to_datetime(out["日期"], errors="coerce")
    out["数量"] = pd.to_numeric(out["数量"], errors="coerce")
    out["单价"] = pd.to_numeric(out["单价"], errors="coerce")

    # 规则 1：关键字段（全部 6 个）缺失的行删除
    out = out.dropna(subset=REQUIRED_COLUMNS)

    # 规则 2：剔除数量 <= 0 的异常行
    out = out[out["数量"] > 0]

    # 规则 3：按全字段（6 个关键字段）去重
    out = out.drop_duplicates(subset=REQUIRED_COLUMNS)

    out = out.reset_index(drop=True)
    if out.empty:
        raise NoValidDataError("清洗后没有任何有效数据，请检查原始文件内容。")
    return out


# ======================= M2：统计聚合 ====================================

def round_amount(value: float) -> float:
    """金额四舍五入到 2 位小数（用 Decimal 避免二进制浮点误差）。"""
    return float(
        Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    )


def add_amount(df: pd.DataFrame) -> pd.DataFrame:
    """增加销售额列：销售额 = 数量 x 单价（保留 2 位小数）。"""
    out = df.copy()
    amount = (out["数量"] * out["单价"]).map(round_amount)
    out[AMOUNT_COLUMN] = amount.map(round_amount)
    return out


def build_salesman_summary(df_with_amount: pd.DataFrame) -> pd.DataFrame:
    """按销售员汇总销售额，销售额降序。"""
    summary = (
        df_with_amount.groupby("销售员", as_index=False)[AMOUNT_COLUMN]
        .sum()
        .sort_values([AMOUNT_COLUMN, "销售员"], ascending=[False, True])
        .reset_index(drop=True)
    )
    summary[AMOUNT_COLUMN] = summary[AMOUNT_COLUMN].map(round_amount)
    return summary


def build_region_summary(df_with_amount: pd.DataFrame) -> pd.DataFrame:
    """按地区汇总销售额，销售额降序。"""
    summary = (
        df_with_amount.groupby("地区", as_index=False)[AMOUNT_COLUMN]
        .sum()
        .sort_values([AMOUNT_COLUMN, "地区"], ascending=[False, True])
        .reset_index(drop=True)
    )
    summary[AMOUNT_COLUMN] = summary[AMOUNT_COLUMN].map(round_amount)
    return summary


def build_monthly_summary(df_with_amount: pd.DataFrame) -> pd.DataFrame:
    """按月汇总销售额（数据可能跨多个月）。"""
    tmp = df_with_amount.copy()
    tmp["月份"] = tmp["日期"].dt.strftime("%Y-%m")
    summary = (
        tmp.groupby("月份", as_index=False)[AMOUNT_COLUMN]
        .sum()
        .sort_values("月份")
        .reset_index(drop=True)
    )
    summary[AMOUNT_COLUMN] = summary[AMOUNT_COLUMN].map(round_amount)
    return summary


def total_amount(df_with_amount: pd.DataFrame) -> float:
    """全部有效数据的销售额合计。"""
    return round_amount(float(df_with_amount[AMOUNT_COLUMN].sum()))


# ======================= M3：Excel 输出 ==================================

SHEET_DETAIL = "明细清洗后"
SHEET_BY_SALESMAN = "按销售员"
SHEET_BY_REGION = "按地区"
SHEET_OVERVIEW = "总览"


def build_overview(
    raw_count: int,
    cleaned_df: pd.DataFrame,
    monthly_summary: pd.DataFrame,
    total: float,
) -> pd.DataFrame:
    """总览表：上半部分是整体指标，下半部分是月度明细（数据可能跨多月）。"""
    cleaned_count = len(cleaned_df)
    rows: list[tuple[str, object]] = [
        ("原始记录数", raw_count),
        ("清洗后记录数", cleaned_count),
        ("剔除记录数", raw_count - cleaned_count),
        ("涉及月份数", int(len(monthly_summary))),
        ("销售总额", total),
        ("", ""),
        ("月度明细", ""),
    ]
    rows += [
        (str(month), float(amount))
        for month, amount in zip(monthly_summary["月份"], monthly_summary[AMOUNT_COLUMN])
    ]
    return pd.DataFrame(rows, columns=["指标", "数值"])


def write_report(
    output_path: str | Path,
    raw_count: int,
    cleaned_with_amount: pd.DataFrame,
    salesman_summary: pd.DataFrame,
    region_summary: pd.DataFrame,
    monthly_summary: pd.DataFrame,
) -> Path:
    """把清洗明细与三张统计表写入 report.xlsx（共 4 个工作表）。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    overview = build_overview(
        raw_count, cleaned_with_amount, monthly_summary, total_amount(cleaned_with_amount)
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        cleaned_with_amount.to_excel(writer, sheet_name=SHEET_DETAIL, index=False)
        salesman_summary.to_excel(writer, sheet_name=SHEET_BY_SALESMAN, index=False)
        region_summary.to_excel(writer, sheet_name=SHEET_BY_REGION, index=False)
        overview.to_excel(writer, sheet_name=SHEET_OVERVIEW, index=False)

    _set_column_widths(output_path)
    return output_path


def _set_column_widths(path: Path) -> None:
    """按内容长度粗略调整列宽，纯粹为了报表可读性。"""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            length = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=8,
            )
            letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[letter].width = min(max(length + 4, 10), 30)
    workbook.save(path)


# ======================= M4：编排 + CLI ==================================

def generate_report(
    input_path: str | Path = "data/sales_raw.xlsx",
    output_path: str | Path = "output/report.xlsx",
) -> Path:
    """一键流程：读取 -> 清洗 -> 统计 -> 输出 report.xlsx。

    原始文件仅以只读方式读取，不做任何写回。
    """
    raw = read_raw(input_path)
    cleaned = clean_data(raw)
    cleaned_with_amount = add_amount(cleaned)

    salesman_summary = build_salesman_summary(cleaned_with_amount)
    region_summary = build_region_summary(cleaned_with_amount)
    monthly_summary = build_monthly_summary(cleaned_with_amount)

    result_path = write_report(
        output_path,
        raw_count=len(raw),
        cleaned_with_amount=cleaned_with_amount,
        salesman_summary=salesman_summary,
        region_summary=region_summary,
        monthly_summary=monthly_summary,
    )

    print(f"原始记录数：{len(raw)}，清洗后记录数：{len(cleaned_with_amount)}，"
          f"剔除：{len(raw) - len(cleaned_with_amount)}")
    for _, row in salesman_summary.iterrows():
        print(f"  [销售员] {row['销售员']}: {row[AMOUNT_COLUMN]}")
    for _, row in region_summary.iterrows():
        print(f"  [地区] {row['地区']}: {row[AMOUNT_COLUMN]}")
    for _, row in monthly_summary.iterrows():
        print(f"  [月份] {row['月份']}: {row[AMOUNT_COLUMN]}")
    print(f"销售总额：{total_amount(cleaned_with_amount)}")
    print(f"报表已生成：{result_path}")
    return result_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Excel 报表生成脚本（考题一）")
    parser.add_argument("--input", default="data/sales_raw.xlsx", help="原始 Excel 路径")
    parser.add_argument("--output", default="output/report.xlsx", help="输出 Excel 路径")
    args = parser.parse_args(argv)

    try:
        generate_report(args.input, args.output)
    except ReportError as exc:
        print(f"[错误] {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
